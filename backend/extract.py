"""Document text extraction.

Pulls per-"page" text out of uploaded vendor documents so the LLM can read them
and so the frontend can display the exact source page behind every matrix cell.

Supported: PDF (pypdf), DOCX (python-docx), XLSX/XLSM (openpyxl), plain text /
markdown / CSV. Scanned/image PDFs come back empty here and are instead sent to
Claude natively via `build_document_content` (vision).
"""
from __future__ import annotations

import base64
import io
from dataclasses import dataclass


@dataclass
class Page:
    file: str
    page: int          # 1-indexed
    text: str


# Word-processing formats have no real pages; we chunk them into pseudo-pages
# of roughly this many characters so provenance stays granular.
_CHUNK_CHARS = 3000


def _chunk(text: str, file: str) -> list[Page]:
    text = text.strip()
    if not text:
        return [Page(file=file, page=1, text="")]
    pages: list[Page] = []
    for i in range(0, len(text), _CHUNK_CHARS):
        pages.append(Page(file=file, page=len(pages) + 1, text=text[i : i + _CHUNK_CHARS]))
    return pages


def _extract_pdf(content: bytes, file: str) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages: list[Page] = []
    for idx, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        pages.append(Page(file=file, page=idx + 1, text=txt.strip()))
    return pages or [Page(file=file, page=1, text="")]


def _extract_docx(content: bytes, file: str) -> list[Page]:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(content))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    # include table cell text too
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _chunk("\n".join(parts), file)


def _extract_xlsx(content: bytes, file: str) -> list[Page]:
    from openpyxl import load_workbook  # reads .xlsx / .xlsm

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows_out: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
            if cells:
                rows_out.append(" | ".join(cells))
        if rows_out:
            parts.append(f"## Sheet: {ws.title}")
            parts.extend(rows_out)
    wb.close()
    return _chunk("\n".join(parts), file)


def extract_pages(filename: str, content: bytes) -> list[Page]:
    """Return a list of Page objects for one uploaded file."""
    name = (filename or "document").strip()
    lower = name.lower()
    try:
        if lower.endswith(".pdf"):
            return _extract_pdf(content, name)
        if lower.endswith(".docx"):
            return _extract_docx(content, name)
        if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
            return _extract_xlsx(content, name)
        if lower.endswith(".xls"):
            return [Page(file=name, page=1, text=f"[Legacy .xls format for {name} could not be read — please re-save as .xlsx or CSV.]")]
        # txt, md, csv, or anything decodable as text
        return _chunk(content.decode("utf-8", errors="replace"), name)
    except Exception as exc:  # pragma: no cover - defensive
        return [Page(file=name, page=1, text=f"[Could not read {name}: {exc}]")]


def pages_to_corpus(pages: list[Page], max_chars: int = 60000) -> tuple[str, bool]:
    """Flatten pages into a single annotated corpus for the LLM.

    Each page is prefixed with a ``[FILE: ... | PAGE n]`` marker so the model can
    cite the exact source. Returns (corpus, truncated).
    """
    blocks: list[str] = []
    total = 0
    truncated = False
    for p in pages:
        body = p.text if p.text else "[no extractable text on this page]"
        block = f"[FILE: {p.file} | PAGE {p.page}]\n{body}"
        if total + len(block) > max_chars:
            truncated = True
            break
        blocks.append(block)
        total += len(block)
    return "\n\n".join(blocks), truncated


def pdf_text_length(content: bytes) -> int:
    """Total extractable text length of a PDF (0 if scanned / image-only)."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        n = 0
        for page in reader.pages:
            try:
                n += len((page.extract_text() or "").strip())
            except Exception:
                pass
        return n
    except Exception:
        return 0


# A PDF with fewer than this many extractable characters is treated as scanned
# and sent to Claude natively (vision) instead of as extracted text.
_SCANNED_PDF_THRESHOLD = 400


def build_document_content(header: str, files: list[tuple[str, bytes]], text: str = "",
                           max_native_bytes: int = 28_000_000) -> list[dict]:
    """Build a Claude message ``content`` list for a set of uploaded documents.

    Text-extractable files (text PDFs, DOCX, TXT) are read as text; PDFs with no
    extractable text (scanned/image-only) are attached as native ``document``
    blocks so Claude reads them with vision — no OCR infrastructure required.
    Document blocks come first, then a single text block. Returns the content list.
    """
    doc_blocks: list[dict] = []
    text_parts: list[str] = []
    native_bytes = 0
    scanned: list[str] = []

    if text and text.strip():
        text_parts.append("PASTED TEXT:\n" + text.strip())

    for filename, content in files:
        lower = (filename or "").lower()
        if lower.endswith(".pdf") and pdf_text_length(content) < _SCANNED_PDF_THRESHOLD:
            if native_bytes + len(content) <= max_native_bytes:
                doc_blocks.append({
                    "type": "document",
                    "source": {"type": "base64", "media_type": "application/pdf",
                               "data": base64.standard_b64encode(content).decode("ascii")},
                    "title": filename,
                })
                native_bytes += len(content)
                scanned.append(filename)
                continue
        # Text path (also covers text-extractable PDFs, DOCX, TXT, CSV, MD).
        pages = extract_pages(filename, content)
        corpus, _ = pages_to_corpus(pages, max_chars=90000)
        if corpus.strip():
            text_parts.append(f"[FILE: {filename}]\n{corpus}")

    body = header
    if text_parts:
        body += "\n\n" + "\n\n".join(text_parts)
    if scanned:
        body += f"\n\n(Attached scanned document(s) — read via vision: {', '.join(scanned)}.)"
    elif not text_parts and not doc_blocks:
        body += "\n\n(No readable content was provided.)"

    content: list[dict] = list(doc_blocks)
    content.append({"type": "text", "text": body})
    return content
